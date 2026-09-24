"""Bureautique : création et analyse de classeurs Excel.

Les fichiers sont écrits dans donnees/fichiers. La création XLSX utilise
un petit générateur OOXML standard, sans dépendance obligatoire. La lecture
utilise openpyxl si disponible, sinon un analyseur XML de secours.
"""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape as xml_echap

from outils import outil

MAX_LIGNES = 5000
MAX_COLONNES = 100
MAX_CELLULES = 20000


def _espace(nom: str):
    from outils.fichiers import _chemin_espace
    return _chemin_espace(nom)


def _xml(valeur: Any) -> str:
    return xml_echap(str(valeur), {'"': "&quot;"})


def _colonne(index: int) -> str:
    lettres = ""
    index += 1
    while index:
        index, reste = divmod(index - 1, 26)
        lettres = chr(65 + reste) + lettres
    return lettres


def _cellule(valeur: Any, ligne: int, colonne: int) -> str:
    reference = f"{_colonne(colonne)}{ligne}"
    if isinstance(valeur, bool):
        return f'<c r="{reference}" t="b"><v>{1 if valeur else 0}</v></c>'
    if isinstance(valeur, (int, float)) and not isinstance(valeur, bool):
        return f'<c r="{reference}"><v>{valeur}</v></c>'
    texte = str(valeur or "")
    if len(texte) > 32767:
        texte = texte[:32767]
    return (f'<c r="{reference}" t="inlineStr"><is><t xml:space="preserve">'
            f'{_xml(texte)}</t></is></c>')


def _parse_contenu(contenu: str) -> list[list[Any]]:
    valeur = (contenu or "").strip()
    if not valeur:
        return [["Colonne 1", "Colonne 2"], ["À compléter", ""]]
    try:
        donnees = json.loads(valeur)
        if isinstance(donnees, dict):
            donnees = donnees.get("donnees", donnees.get("lignes", []))
        if isinstance(donnees, list) and donnees and isinstance(donnees[0], dict):
            donnees = [list(donnees[0].keys())] + [list(row.values()) for row in donnees]
    except json.JSONDecodeError:
        donnees = [ligne.split("\t") for ligne in valeur.splitlines()]
    if not isinstance(donnees, list) or not donnees:
        donnees = [["Colonne 1"], [""]]
    lignes: list[list[Any]] = []
    for ligne in donnees[:MAX_LIGNES]:
        if isinstance(ligne, (list, tuple)):
            lignes.append(list(ligne)[:MAX_COLONNES])
        else:
            lignes.append([ligne][:MAX_COLONNES])
    largeur = max((len(ligne) for ligne in lignes), default=1)
    return [ligne + [""] * (largeur - len(ligne)) for ligne in lignes]


def _xlsx(lignes: list[list[Any]], feuille: str) -> bytes:
    lignes = lignes[:MAX_LIGNES]
    largeur = min(max((len(ligne) for ligne in lignes), default=1), MAX_COLONNES)
    lignes = [ligne + [""] * (largeur - len(ligne)) for ligne in lignes]
    lignes_xml = []
    for numero, ligne in enumerate(lignes, 1):
        cellules = "".join(_cellule(valeur, numero, colonne)
                          for colonne, valeur in enumerate(ligne[:largeur]))
        lignes_xml.append(f'<row r="{numero}">{cellules}</row>')
    feuille_xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                   f'<dimension ref="A1:{_colonne(max(0, largeur - 1))}{len(lignes)}"/>'
                   '<sheetViews><sheetView workbookViewId="0"/></sheetViews>'
                   '<sheetFormatPr defaultRowHeight="15"/>'
                   '<cols><col min="1" max="%d" width="18" customWidth="1"/></cols>'
                   '<sheetData>%s</sheetData></worksheet>' % (largeur, "".join(lignes_xml)))
    content_types = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                     '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                     '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                     '<Default Extension="xml" ContentType="application/xml"/>'
                     '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                     '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                     '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
                     '</Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            '</Relationships>')
    workbook = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                f'<sheets><sheet name="{_xml(feuille[:31] or "Feuille1")}" sheetId="1" r:id="rId1"/></sheets>'
                '</workbook>')
    workbook_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                     '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                     '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
                     '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
                     '</Relationships>')
    styles = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
              '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
              '<fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
              '<borders count="1"><border/></borders>'
              '<cellStyleXfs count="1"><xf/></cellStyleXfs>'
              '<cellXfs count="1"><xf xfId="0"/></cellXfs>'
              '</styleSheet>')
    from io import BytesIO
    tampon = BytesIO()
    with zipfile.ZipFile(tampon, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        archive.writestr("xl/styles.xml", styles)
        archive.writestr("xl/worksheets/sheet1.xml", feuille_xml)
    return tampon.getvalue()


@outil("creer_excel",
       "Crée un vrai classeur Excel .xlsx dans l'espace de travail. `contenu` "
       "est un tableau JSON de lignes ou un texte séparé par tabulations.",
       {"nom": {"type": "str", "obligatoire": True, "description": "ex. budget.xlsx"},
        "contenu": {"type": "str", "obligatoire": True,
                    "description": "JSON [[\"Dépense\",\"Montant\"],[\"Loyer\",850]]"},
        "feuille": {"type": "str", "obligatoire": False,
                    "description": "nom de la feuille (défaut Feuille1)"}},
       categorie="documents", risque="moyen",
       exemple='{"outil": "creer_excel", "parametres": {"nom": "budget.xlsx", "contenu": "[[\"Poste\",\"Montant\"],[\"Loyer\",850]]"}}')
def creer_excel(nom: str, contenu: str, feuille: str = "Feuille1") -> str:
    nom = nom.strip()
    if not nom.lower().endswith(".xlsx"):
        nom += ".xlsx"
    chemin = _espace(nom)
    lignes = _parse_contenu(contenu)
    chemin.parent.mkdir(parents=True, exist_ok=True)
    chemin.write_bytes(_xlsx(lignes, feuille))
    return f"Excel créé : {chemin.name} ({len(lignes)} ligne(s), {chemin.stat().st_size} octets)."


def _analyser_xlsx_zip(chemin: Path) -> dict:
    with zipfile.ZipFile(chemin) as archive:
        noms = set(archive.namelist())
        if "xl/worksheets/sheet1.xml" not in noms:
            raise ValueError("classeur sans feuille lisible")
        xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8", "replace")
    lignes = []
    for brut in re.findall(r"<row[^>]*>(.*?)</row>", xml, re.S):
        cellules = []
        for valeur in re.findall(r"<c[^>]*?(?:t=\"([^\"]+)\")?[^>]*>(.*?)</c>", brut, re.S):
            type_cellule, contenu = valeur
            texte = re.sub(r"<[^>]+>", "", contenu)
            cellules.append(texte)
        lignes.append(cellules)
    return {"feuilles": 1, "lignes": len(lignes),
            "colonnes": max((len(l) for l in lignes), default=0),
            "formules": xml.count("<f"), "non_vides": sum(1 for l in lignes for c in l if c),
            "apercu": lignes[:5]}


@outil("analyser_excel",
       "Analyse un classeur Excel (.xlsx) : feuilles, dimensions, formules, "
       "valeurs non vides et aperçu des premières lignes. Lecture seule.",
       {"chemin": {"type": "str", "obligatoire": True,
                   "description": "classeur dans donnees/fichiers ou chemin absolu"}},
       categorie="documents", risque="eleve",
       exemple='{"outil": "analyser_excel", "parametres": {"chemin": "budget.xlsx"}}')
def analyser_excel(chemin: str) -> str:
    if not chemin:
        return "Chemin manquant."
    try:
        cible = _espace(chemin) if not (Path(chemin).is_absolute() or ":" in chemin or "\\" in chemin) else Path(chemin).expanduser().resolve()
    except ValueError as e:
        return str(e)
    if not cible.is_file():
        return f"Fichier introuvable : {chemin}"
    from outils.analyse import _refuse
    if _refuse(cible):
        return "Analyse refusée : ce fichier ressemble à un secret."
    try:
        import openpyxl  # type: ignore[import-not-found]
        classeur = openpyxl.load_workbook(cible, read_only=True, data_only=False)
        feuilles = []
        apercu = []
        for feuille in classeur.worksheets:
            lignes = list(feuille.iter_rows(max_row=5, values_only=True))
            feuilles.append(feuille.title)
            if not apercu:
                apercu = lignes
        return (f"Excel : {cible.name}\nFeuille(s) : {', '.join(feuilles)}\n"
                f"Dimensions : {classeur.active.max_row} lignes × {classeur.active.max_column} colonnes\n"
                f"Formules : {sum(1 for row in classeur.active.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('='))}\n"
                f" Aperçu : {apercu}")
    except ImportError:
        try:
            resultat = _analyser_xlsx_zip(cible)
            return (f"Excel : {cible.name}\nFeuille(s) : {resultat['feuilles']}\n"
                    f"Dimensions : {resultat['lignes']} lignes × {resultat['colonnes']} colonnes\n"
                    f"Formules : {resultat['formules']} · valeurs non vides : {resultat['non_vides']}\n"
                    f"Aperçu : {resultat['apercu']}")
        except Exception as e:
            return f"Classeur illisible : {e}"
    except Exception as e:
        return f"Analyse Excel impossible : {str(e)[:180]}"


@outil("analyser_documents",
       "Analyse un dossier de documents (Word, PDF, Excel, texte, CSV, JSON, "
       "images) et renvoie un résumé de chaque fichier. Lecture seule.",
       {"chemin": {"type": "str", "obligatoire": False,
                   "description": "dossier ; vide = donnees/fichiers"},
        "limite": {"type": "int", "obligatoire": False,
                    "description": "maximum de documents (défaut 20)"}},
       categorie="documents", risque="eleve",
       exemple='{"outil": "analyser_documents", "parametres": {"chemin": "C:/Users/moi/Documents"}}')
def analyser_documents(chemin: str = "", limite: int = 20) -> str:
    if not chemin:
        racine = Path(_espace("."))
    else:
        try:
            racine = _espace(chemin) if not (Path(chemin).is_absolute() or ":" in chemin or "\\" in chemin) else Path(chemin).expanduser().resolve()
        except ValueError as e:
            return str(e)
    if not racine.is_dir():
        return f"Dossier introuvable : {racine}"
    from outils.analyse import _resume, _refuse
    extensions = {".docx", ".pdf", ".xlsx", ".txt", ".md", ".csv", ".json", ".png", ".jpg", ".jpeg"}
    limite = max(1, min(int(limite), 100))
    resumes = []
    for fichier in sorted(racine.rglob("*")):
        if len(resumes) >= limite:
            break
        if not fichier.is_file() or fichier.suffix.lower() not in extensions or _refuse(fichier):
            continue
        if fichier.stat().st_size > 25 * 1024 * 1024:
            resumes.append(f"- {fichier.name} : trop volumineux")
        elif fichier.suffix.lower() == ".xlsx":
            resumes.append(f"- {fichier.name} : {analyser_excel(str(fichier))[:700]}")
        else:
            resumes.append(f"- {fichier.name} : {_resume(fichier)[:700]}")
    return "Analyse de documents :\n" + ("\n".join(resumes) if resumes else "Aucun document trouvé.")
